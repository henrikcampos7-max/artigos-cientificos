from __future__ import annotations

import argparse
import csv
import re
import unicodedata
from pathlib import Path

from docx import Document


def normalize(value: str) -> str:
    value = unicodedata.normalize("NFKD", value)
    value = "".join(character for character in value if not unicodedata.combining(character))
    return re.sub(r"\s+", " ", value).strip().lower()


def paragraph_texts(document: Document) -> list[str]:
    values = [re.sub(r"\s+", " ", paragraph.text).strip() for paragraph in document.paragraphs]
    return [value for value in values if value]


def reference_block(paragraphs: list[str]) -> list[str]:
    start = None
    for index, paragraph in enumerate(paragraphs):
        if re.fullmatch(r"(?:[ivxlcdm]+\.?\s*)?referencias(?: bibliograficas)?", normalize(paragraph)):
            start = index + 1
    if start is None:
        return []
    candidates = paragraphs[start:]
    return [
        paragraph
        for paragraph in candidates
        if len(paragraph) >= 30 and re.search(r"\b(?:19|20)\d{2}[a-z]?\b", paragraph)
    ]


def audit(path: Path, root: Path) -> dict[str, object]:
    document = Document(path)
    paragraphs = paragraph_texts(document)
    full_text = "\n".join(paragraphs)
    references = reference_block(paragraphs)
    reference_text = "\n".join(references)
    doi_values = {
        value.lower().rstrip(".,;)")
        for value in re.findall(r"10\.\d{4,9}/[-._;()/:A-Z0-9]+", reference_text, re.I)
    }
    recent_references = [
        reference for reference in references if re.search(r"\b202[1-6][a-z]?\b", reference)
    ]
    title = next((paragraph for paragraph in paragraphs[:12] if len(paragraph) >= 12), "")
    properties = document.core_properties
    return {
        "arquivo": path.relative_to(root).as_posix(),
        "titulo_estimado": title,
        "status": "rascunho_nao_validado",
        "palavras": len(re.findall(r"\b\w+[\w-]*\b", full_text, re.UNICODE)),
        "referencias_estimadas": len(references),
        "dois_unicos": len(doi_values),
        "referencias_2021_2026_estimadas": len(recent_references),
        "proporcao_recente_estimada": (
            round(len(recent_references) / len(references), 3) if references else ""
        ),
        "menciona_limitacoes": bool(re.search(r"\blimita[cç][aã]o|\blimita[cç][oõ]es\b", full_text, re.I)),
        "afirma_revisao_narrativa": bool(re.search(r"\brevis[aã]o narrativa\b", full_text, re.I)),
        "menciona_revisao_sistematica": bool(re.search(r"\brevis[aã]o sistem[aá]tica\b", full_text, re.I)),
        "tabelas": len(document.tables),
        "autor_metadado": properties.author or "",
        "criado_metadado": properties.created.isoformat() if properties.created else "",
        "observacao": "Triagem automática; exige leitura integral e revisão humana.",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Audita DOCX sem modificar os arquivos-fonte.")
    parser.add_argument("root", type=Path, help="Pasta que contém os DOCX")
    parser.add_argument("--csv", type=Path, required=True, help="CSV de saída")
    args = parser.parse_args()
    root = args.root.resolve()
    files = sorted(path for path in root.rglob("*.docx") if not path.name.startswith("~$"))
    rows = []
    for path in files:
        try:
            rows.append(audit(path, root))
        except Exception as exc:  # registra bloqueios/corrupção sem interromper o lote
            rows.append(
                {
                    "arquivo": path.relative_to(root).as_posix(),
                    "status": "erro_de_leitura",
                    "observacao": f"{type(exc).__name__}: {exc}",
                }
            )
    fieldnames = [
        "arquivo",
        "titulo_estimado",
        "status",
        "palavras",
        "referencias_estimadas",
        "dois_unicos",
        "referencias_2021_2026_estimadas",
        "proporcao_recente_estimada",
        "menciona_limitacoes",
        "afirma_revisao_narrativa",
        "menciona_revisao_sistematica",
        "tabelas",
        "autor_metadado",
        "criado_metadado",
        "observacao",
    ]
    args.csv.parent.mkdir(parents=True, exist_ok=True)
    with args.csv.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    print(f"{len(rows)} arquivo(s) auditado(s); saída: {args.csv}")


if __name__ == "__main__":
    main()
